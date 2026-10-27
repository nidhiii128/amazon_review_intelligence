"""
Builds the ReviewIQ presentation. Each slide is defined once as data (a list
of elements in inches); the same spec is rendered to a PPTX (python-pptx) and
to an HTML preview that is screenshotted to check every slide for overflow.

Usage: python3 build_deck.py <out.pptx> <preview.html>
"""
import html
import pathlib
import re
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt
from pptx.oxml.ns import qn
from lxml import etree

BASE = pathlib.Path(__file__).resolve().parent.parent
ASSETS = BASE / "assets"
ROOT = BASE.parent
W, H = 13.333, 7.5
FONT = "Arial"

NAVY, NAVY2, ACCENT, BLUE, TEAL = "131A2B", "1F2A44", "E8622C", "2563EB", "0F766E"
INK, MUTED, LINE, PAPER, WHITE = "1F2937", "5B6475", "E4E1DA", "FAF9F6", "FFFFFF"
SOFT_ORG, SOFT_BLUE, SOFT_TEAL = "FFF4EC", "EFF6FF", "F0FDFA"

TEAM = [("Nidhi Bhawari", "23102A0006"), ("Aarti Sakpal", "23102A0007"),
        ("Gayatri Patil", "23102A0016"), ("Prachi Wakode", "23102A0020")]
GUIDE = "Prof. Sneha Annappanavar"
COLLEGE = "Vidyalankar Institute of Technology"


def img(path, x, y, w=None, h=None):
    return {"t": "img", "path": str(path), "x": x, "y": y, "w": w, "h": h}


def rect(x, y, w, h, fill=WHITE, line=LINE, radius=True, lw=1.0):
    return {"t": "rect", "x": x, "y": y, "w": w, "h": h, "fill": fill, "line": line, "radius": radius, "lw": lw}


def text(x, y, w, h, paras, anchor="top", pad=0.08):
    """paras: list of dicts {s, size, bold, color, bullet, align, space}"""
    return {"t": "text", "x": x, "y": y, "w": w, "h": h, "paras": paras, "anchor": anchor, "pad": pad}


def P(s, size=20, bold=False, color=INK, bullet=False, align="l", space=6):
    return {"s": s, "size": size, "bold": bold, "color": color, "bullet": bullet, "align": align, "space": space}


def arrow(x, y, w, h, fill=ACCENT):
    return {"t": "arrow", "x": x, "y": y, "w": w, "h": h, "fill": fill}


def table(x, y, w, colw, rows, size=16, hl=(), best=(), rowh=0.42):
    return {"t": "table", "x": x, "y": y, "w": w, "colw": colw, "rows": rows, "size": size, "hl": hl, "best": best, "rowh": rowh}


def stat(x, y, w, h, big, small, color=ACCENT, fill=SOFT_ORG):
    return [rect(x, y, w, h, fill=fill, line=LINE),
            text(x, y + 0.08, w, h * 0.55, [P(big, 34, True, color, align="c", space=0)], "middle"),
            text(x + 0.1, y + h * 0.55, w - 0.2, h * 0.42, [P(small, 13, False, MUTED, align="c", space=0)], "top")]


def card(x, y, w, h, title, paras, accent=ACCENT, fill=WHITE, tsize=20):
    els = [rect(x, y, w, h, fill=fill), rect(x, y + 0.12, 0.09, 0.42, fill=accent, line=accent, radius=False)]
    els.append(text(x + 0.22, y + 0.08, w - 0.3, 0.55, [P(title, tsize, True, NAVY, space=0)], "middle"))
    els.append(text(x + 0.14, y + 0.68, w - 0.28, h - 0.76, paras))
    return els


def callout(x, y, w, h, s, size=17):
    return [rect(x, y, w, h, fill=NAVY, line=NAVY), text(x + 0.15, y, w - 0.3, h, [P(s, size, False, WHITE, space=0)], "middle")]


SLIDES = []


def slide(title=None, notes="", dark=False):
    s = {"title": title, "notes": notes, "dark": dark, "els": []}
    SLIDES.append(s)
    return s["els"]


# =============================================================== 1 TITLE
e = slide(dark=True, notes="Introduce the project, the team and the guide. ReviewIQ turns thousands of unstructured product reviews into answers for buyers and sellers.")
e += [rect(0, 0, W, H, fill=NAVY, line=NAVY, radius=False), rect(0, 5.05, W, 0.09, fill=ACCENT, line=ACCENT, radius=False)]
e += [text(0.7, 0.45, 11.5, 0.5, [P(COLLEGE.upper(), 18, True, "C8D2E6")])]
e += [text(0.7, 1.0, 11.5, 1.3, [P("ReviewIQ", 66, True, WHITE, space=0)]),
      text(0.7, 2.25, 11.9, 1.0, [P("AI-Based Product Review Intelligence and Consumer Insight System", 25, False, "DBE3F3", space=0)]),
      text(0.7, 3.35, 11.5, 0.6, [P("NLP  |  Deep Learning  |  Advanced Machine Learning", 20, True, "FFB48F", space=0)]),
      text(0.7, 3.95, 11.5, 0.5, [P("Project Presentation", 18, False, "C8D2E6", space=0)])]
e += [text(0.7, 5.35, 6.0, 0.4, [P("PROJECT TEAM", 14, True, "FFB48F", space=0)])]
for i, (n, r) in enumerate(TEAM):
    e += [text(0.7, 5.75 + i * 0.36, 3.4, 0.36, [P(n, 17, True, WHITE, space=0)], "middle"),
          text(4.0, 5.75 + i * 0.36, 2.6, 0.36, [P(r, 17, False, "C8D2E6", space=0)], "middle")]
e += [text(8.2, 5.35, 4.6, 0.4, [P("UNDER THE GUIDANCE OF", 14, True, "FFB48F", space=0)]),
      text(8.2, 5.75, 4.8, 0.5, [P(GUIDE, 22, True, WHITE, space=0)], "middle"),
      text(8.2, 6.25, 4.8, 0.4, [P(COLLEGE, 15, False, "C8D2E6", space=0)])]

# =============================================================== 2 WHY
e = slide("Why this project?", "A popular product can have thousands of reviews. The star average hides what customers actually say. We built one system that answers three different questions.")
e += [text(0.7, 1.4, 6.2, 5.2, [
    P("**Reviews are abundant but unstructured.** A single popular product can attract thousands of free-text opinions.", 19, bullet=True, space=9),
    P("**The star average hides the detail.** It cannot say that the sound is great and the battery is poor.", 19, bullet=True, space=9),
    P("**Nobody can read everything.** Automated analysis is central to e-commerce.", 19, bullet=True, space=9),
    P("**Related tasks are usually studied separately:** sentiment, aspects, helpfulness, spam. We combine them on one corpus.", 19, bullet=True, space=9)])]
qs = [("Should I buy this?", "Buyers: strengths, weaknesses and evidence, per product", BLUE, SOFT_BLUE),
      ("What should we improve?", "Sellers: ranked complaint areas across the category", ACCENT, SOFT_ORG),
      ("Can I trust these reviews?", "Review quality: helpfulness and suspicious-review signals", TEAL, SOFT_TEAL)]
for i, (q, d, c, f) in enumerate(qs):
    y = 1.45 + i * 1.75
    e += [rect(7.4, y, 5.3, 1.55, fill=f, line=LINE), rect(7.4, y + 0.15, 0.1, 1.25, fill=c, line=c, radius=False),
          text(7.7, y + 0.1, 4.9, 0.6, [P(q, 21, True, c, space=0)], "middle"),
          text(7.7, y + 0.7, 4.9, 0.8, [P(d, 16, False, MUTED, space=0)])]

# =============================================================== 3 OBJECTIVES
e = slide("Objectives and scope", "Seven objectives. The scope is deliberately bounded: one category, a 50,000 review sample, labels from star ratings, and an unsupervised trust score that makes no accuracy claim.")
e += card(0.6, 1.4, 7.4, 5.3, "Objectives", [
    P("Prepare a clean, labelled corpus enriched with product names, brands and images", 18, bullet=True),
    P("Classify sentiment (positive, neutral, negative) with NLP, deep learning and advanced ML models and compare them fairly", 18, bullet=True),
    P("Build an LSTM classifier, report its metrics and visualise each sentiment with word clouds", 18, bullet=True),
    P("Extract aspect-level sentiment, topics and clusters", 18, bullet=True),
    P("Predict helpfulness and flag suspicious reviews", 18, bullet=True),
    P("Provide semantic search and grounded question answering with evidence", 18, bullet=True),
    P("Deliver everything in one interactive dashboard", 18, bullet=True)])
e += card(8.3, 1.4, 4.45, 5.3, "Scope and boundaries", [
    P("One product category: **Amazon Electronics**", 17, bullet=True),
    P("English reviews, **50,000-review** working sample", 17, bullet=True),
    P("Sentiment labels come from **star ratings**", 17, bullet=True),
    P("Suspicious-review detection is **unsupervised**: no fake-review accuracy is claimed", 17, bullet=True),
    P("Answers come only from the selected product's own reviews", 17, bullet=True)], accent=BLUE)

# =============================================================== 4 DATASET
e = slide("Dataset and preprocessing", "Amazon Reviews 2023, Electronics. We cleaned, de-duplicated and labelled the reviews, and joined product metadata on the parent ASIN. The sample is a contiguous block of a reviewer-grouped file, which we state as a limitation.")
for i, (b, s, c, f) in enumerate([("50,000", "reviews", ACCENT, SOFT_ORG), ("34,365", "products", BLUE, SOFT_BLUE),
                                  ("8,883", "reviewers", TEAL, SOFT_TEAL), ("80 / 7 / 13", "% positive / neutral / negative", ACCENT, SOFT_ORG)]):
    e += stat(0.6 + i * 3.1, 1.4, 2.9, 1.55, b, s, c, f)
e += card(0.6, 3.2, 12.15, 3.5, "How the data was prepared", [
    P("**Source:** Amazon Reviews 2023 (McAuley Lab), Electronics category", 17, bullet=True),
    P("**Cleaning:** removed missing ratings and very short texts, stripped HTML and URLs, dropped duplicate (reviewer, product, text) records", 17, bullet=True),
    P("**Labels:** stars 4 and 5 positive, 3 neutral, 1 and 2 negative. Weak labels, so macro-F1 is the main metric", 17, bullet=True),
    P("**Metadata:** title, brand and image joined on the parent ASIN (34,364 of 34,365 products matched)", 17, bullet=True),
    P("**Evaluation:** one shared 40,000 train and 10,000 test split for every model", 17, bullet=True)])

# =============================================================== 5 ARCHITECTURE
e = slide("System architecture", "Reviews and metadata are cleaned and joined, then three parallel branches feed product intelligence and a grounded question answering layer, all in one dashboard.")
e += [img(ASSETS / "figs" / "fig_framework.png", 1.37, 1.3, w=10.6)]
e += callout(0.9, 6.0, 11.5, 0.75, "Three pillars on one shared corpus: NLP, deep learning and advanced machine learning, fused into product-level intelligence.", 18)

# =============================================================== 6 NLP
e = slide("NLP branch", "Five NLP components. The lexicon and TF-IDF models are baselines. LDA discovers topics. The aspect module keeps the matched sentences as evidence.")
items = [("VADER", "Rule-based sentiment lexicon. Zero-training baseline.", BLUE),
         ("TF-IDF + LR / SVM", "Unigrams and bigrams, 30,000 features, balanced class weights.", BLUE),
         ("LDA topics", "8 topics found without labels: audio, storage, cables, cameras, cases and more.", BLUE),
         ("Aspect sentiment", "10 aspects (battery, sound, camera ...). Matched sentences are kept as evidence.", BLUE),
         ("MiniLM embeddings", "384-d sentence vectors power semantic search and question answering.", BLUE)]
for i, (t, d, c) in enumerate(items):
    col, row = i % 3, i // 3
    e += card(0.6 + col * 4.1, 1.4 + row * 2.75, 3.9, 2.55, t, [P(d, 17)], accent=c)
e += [rect(8.8, 4.15, 3.9, 2.55, fill=SOFT_BLUE, line=LINE),
      text(9.0, 4.25, 3.5, 2.35, [P("Why it matters", 19, True, BLUE),
                                  P("The same corpus supports both **review-level labels** and **aspect-level insight**.", 16, space=8),
                                  P("Aspect results explain why a product is rated the way it is.", 16)], "middle")]

# =============================================================== 7 DEEP LEARNING
e = slide("Deep learning: BiLSTM with attention", "We trained a bidirectional LSTM with attention from scratch in PyTorch. DistilBERT was fine-tuned on a small CPU budget. Attention lets the model weight the words that matter.")
e += [rect(0.6, 1.4, 4.0, 5.3, fill=WHITE), img(ASSETS / "poster_figs" / "poster_lstm.png", 0.8, 1.55, w=3.6)]
e += card(4.9, 1.4, 7.85, 3.15, "BiLSTM + Attention (from scratch)", [
    P("Vocabulary of **14,937** words, 128-d embeddings, **2 x 128** bidirectional LSTM", 17, bullet=True),
    P("Attention pooling learns which tokens matter most", 17, bullet=True),
    P("Class-weighted loss, early stopping on validation **macro-F1**", 17, bullet=True),
    P("**2.57 M** parameters, about **7 minutes** on a laptop GPU", 17, bullet=True)], accent=ACCENT)
e += card(4.9, 4.75, 7.85, 1.95, "DistilBERT (fine-tuned)", [
    P("Pre-trained transformer, fine-tuned on **6,000** reviews on a CPU (about 60 min)", 16, bullet=True),
    P("Best raw accuracy, but one-sixth of the LSTM's training data", 16, bullet=True)], accent=BLUE)

# =============================================================== 8 ADVANCED ML
e = slide("Advanced machine learning", "XGBoost predicts helpfulness. KMeans is a negative finding that we report honestly. The Isolation Forest is unsupervised and we fixed its features after finding they were backwards.")
e += card(0.6, 1.4, 3.95, 5.3, "XGBoost helpfulness", [
    P("Target: at least one helpful vote", 16, bullet=True),
    P("10 engineered features, incl. aspects mentioned (from ABSA)", 16, bullet=True),
    P("15-candidate randomized search, 3-fold CV", 16, bullet=True),
    P("**70.5%** accuracy, helpful-class F1 **0.555**", 16, bullet=True),
    P("Length, aspect count and rating matter most", 16, bullet=True)], accent=TEAL)
e += card(4.7, 1.4, 3.95, 5.3, "MiniBatchKMeans", [
    P("k = 6 clusters on a separate TF-IDF space", 16, bullet=True),
    P("Only **2** clusters are clearly topical (bags and cases; sound and headphones)", 16, bullet=True),
    P("The other 4 are generic praise (\"works great\", \"five stars\")", 16, bullet=True),
    P("A negative finding: LDA is the better topic tool on short, rating-skewed text", 16, bullet=True)], accent=TEAL)
e += card(8.8, 1.4, 3.95, 5.3, "Isolation Forest", [
    P("Unsupervised: **no fake-review accuracy** is claimed", 16, bullet=True),
    P("Features: very short text, caps, exclamations, verified, account volume, near-duplicate similarity", 16, bullet=True),
    P("Flags **4,999** reviews (10%), 36 near-duplicates", 16, bullet=True),
    P("Early version flagged long, negative reviews. Fixed with directional features", 16, bullet=True)], accent=TEAL)

# =============================================================== 9 RAG
e = slide("Semantic search and grounded Q and A", "The question is embedded, only the selected product's reviews are searched, and if the best similarity is below 0.30 the system says it does not have enough information instead of letting the generator guess.")
boxes = [("Question", "e.g. \"Is it good for cutting cable TV?\"", BLUE, SOFT_BLUE), ("Embed + retrieve", "MiniLM, top 5 reviews of this product", BLUE, SOFT_BLUE),
         ("Similarity floor", "best similarity at least 0.30?", ACCENT, SOFT_ORG), ("Answer + evidence", "FLAN-T5 answer, retrieved reviews shown", TEAL, SOFT_TEAL)]
for i, (t, d, c, f) in enumerate(boxes):
    x = 0.6 + i * 3.2
    e += [rect(x, 1.55, 2.75, 1.9, fill=f, line=LINE), text(x + 0.1, 1.6, 2.55, 0.65, [P(t, 17, True, c, align="c", space=0)], "middle"),
          text(x + 0.12, 2.3, 2.5, 1.1, [P(d, 14, False, MUTED, align="c", space=0)], "top")]
    if i < 3: e += [arrow(x + 2.78, 2.25, 0.36, 0.4)]
e += card(0.6, 3.75, 6.0, 2.95, "Why a similarity floor?", [
    P("Small generators do not reliably follow \"say you do not know\"", 17, bullet=True),
    P("Before the floor, a warranty question got a confident \"No\" at similarity 0.25", 17, bullet=True),
    P("Weak evidence now returns: **not enough information**", 17, bullet=True)], accent=ACCENT)
e += card(6.85, 3.75, 5.9, 2.95, "Measured effect", [
    P("**97.8%** of answerable questions accepted", 17, bullet=True),
    P("**95.6%** of plausible but unanswerable questions declined", 17, bullet=True),
    P("**100%** of off-topic questions declined", 17, bullet=True),
    P("Retrieval beats TF-IDF on paraphrased queries: **0.54** vs 0.32", 17, bullet=True)], accent=TEAL)

# =============================================================== 10 RESULTS
e = slide("Results: sentiment classification", "All five models are scored on the same 10,000 test reviews. The BiLSTM has the best macro-F1 and DistilBERT the best accuracy. The BiLSTM beats the SVM significantly on macro-F1, but not DistilBERT.")
e += [table(0.6, 1.4, 6.6, [2.45, 1.3, 1.4, 1.45],
            [["Model", "Accuracy", "Macro-F1", "F1 neutral"], ["VADER", ".813", ".512", ".080"], ["TF-IDF + LR", ".860", ".689", ".400"],
             ["TF-IDF + SVM", ".888", ".693", ".382"], ["BiLSTM + Attention", ".888", ".718", ".446"], ["DistilBERT", ".899", ".702", ".368"]],
            size=16, hl=(4, 5), best={(4, 2), (4, 3), (5, 1)}, rowh=0.5)]
e += callout(0.6, 4.65, 6.6, 1.95, "BiLSTM minus SVM: +0.025 macro-F1 (95% CI +0.010 to +0.039), significant. Accuracy is identical. The gain is on the neutral class.", 17)
e += [rect(7.5, 1.4, 5.25, 5.2, fill=WHITE), img(ASSETS / "poster_figs" / "poster_ci.png", 7.6, 1.55, w=5.05),
      text(7.6, 4.0, 5.05, 2.5, [P("95% bootstrap intervals, 2,000 resamples of the test set.", 14, False, MUTED, space=6),
                                 P("Neutral is the hardest class for every model: it is 7% of the data and sits between positive and negative.", 15, False, INK)], "top")]

# =============================================================== 11 ABLATION
e = slide("Ablation and cost", "Three variants times three seeds on the same split. Class weighting, not attention, drives most of the macro-F1 gain, and an unweighted LSTM matches DistilBERT accuracy. The LSTM is far cheaper to run.")
e += [text(0.6, 1.35, 7.0, 0.5, [P("What drives the LSTM's gain? (mean of 3 seeds)", 18, True, NAVY, space=0)], "middle"),
      table(0.6, 1.9, 7.0, [3.9, 1.5, 1.6],
            [["Variant", "Accuracy", "Macro-F1"], ["Full (attention + class weights)", ".891", ".722"], ["No class weights", ".899", ".709"],
             ["Mean pooling, no attention", ".889", ".715"], ["DistilBERT (reference, 1 run)", ".899", ".702"]],
            size=16, hl=(1,), best={(1, 2), (2, 1)}, rowh=0.5)]
e += callout(0.6, 4.65, 7.0, 1.95, "Class weights add +0.012 macro-F1 and cost 0.007 accuracy, in every seed. Attention adds only about +0.007. Seed-to-seed spread is about 0.004.", 17)
e += stat(8.0, 1.4, 4.75, 1.55, "26x", "fewer parameters than DistilBERT (2.57 M vs 66.96 M)", ACCENT, SOFT_ORG)
e += stat(8.0, 3.15, 4.75, 1.55, "57x", "faster on CPU (975 vs 17 reviews per second)", BLUE, SOFT_BLUE)
e += stat(8.0, 4.9, 4.75, 1.55, "13x", "faster on GPU (5,688 vs 451 reviews per second)", TEAL, SOFT_TEAL)

# =============================================================== 12 ASPECTS
e = slide("Aspects and the language of sentiment", "Battery is the top complaint, then connectivity and customer service. The word clouds weight words by how characteristic they are of each class.")
e += [rect(0.6, 1.4, 6.3, 5.3, fill=WHITE), img(ASSETS / "poster_figs" / "poster_aspects.png", 0.75, 1.55, w=6.0)]
e += callout(0.75, 5.35, 6.0, 1.2, "Battery: 20.2% negative mentions. Connectivity and customer service: 19.3% each. Build quality: 8.1%.", 16)
e += [rect(7.2, 1.4, 5.55, 5.3, fill=WHITE),
      text(7.35, 1.45, 5.3, 0.45, [P("What each sentiment sounds like", 19, True, NAVY, space=0)], "middle"),
      img(ROOT / "reports" / "figures" / "wordcloud_negative.png", 7.35, 1.95, w=2.6), img(ROOT / "reports" / "figures" / "wordcloud_positive.png", 10.05, 1.95, w=2.6),
      text(7.35, 3.5, 2.6, 0.35, [P("Negative", 14, True, "9A3412", align="c", space=0)]), text(10.05, 3.5, 2.6, 0.35, [P("Positive", 14, True, "166534", align="c", space=0)]),
      text(7.35, 3.95, 5.3, 2.6, [P("Negative: \"return\", \"waste\", \"money\", \"stopped\", \"broke\"", 16, bullet=True),
                                  P("Positive: \"easy\", \"love\", \"nice\", \"perfect\"", 16, bullet=True),
                                  P("Neutral: hedged words such as \"okay\", \"fine\", \"better\"", 16, bullet=True)])]

# =============================================================== 13 TRUST
e = slide("Review quality, trust and error analysis", "Helpfulness is moderately predictable. The suspicious-review score is unsupervised. Our error analysis shows that some model errors are really label noise.")
e += card(0.6, 1.4, 3.95, 4.2, "Helpfulness (XGBoost)", [
    P("**70.5%** accuracy, precision 0.458, recall 0.702, F1 0.555", 16, bullet=True),
    P("Modest because helpfulness depends on exposure, which no review feature captures", 16, bullet=True),
    P("The aspect count from the ABSA module ranks second: one component feeds another", 16, bullet=True)], accent=TEAL, tsize=18)
e += card(4.7, 1.4, 3.95, 4.2, "Suspicious reviews", [
    P("**4,999** flagged (10%), **36** near-duplicates", 16, bullet=True),
    P("Dashboard: threshold slider and a table of flag, score, review and signals", 16, bullet=True),
    P("Worded as \"potentially suspicious based on automated pattern analysis\". Never \"fake\"", 16, bullet=True)], accent=ACCENT, tsize=18)
e += card(8.8, 1.4, 3.95, 4.2, "Label noise in errors", [
    P("155 of 1,116 test errors are confident", 15, bullet=True),
    P("**5 stars, predicted negative:** \"Piece of junk !!! Poor quality, high price .. it lasted 10weeks\"", 15, bullet=True),
    P("**1 star, predicted positive:** \"Worked out of the box. love it\"", 15, bullet=True),
    P("The stars contradict the text, so part of the error is label noise", 15, bullet=True)], accent=BLUE, tsize=18)
e += callout(0.6, 5.85, 12.15, 0.9, "Takeaway: accuracy here measures agreement with star ratings. Some errors are label noise, and flagged reviews are anomalies, not proven fakes.", 17)

# =============================================================== 14 DASH 1
e = slide("The dashboard (1 of 2)", "Nine pages. The Home page shows the three pillars. Product Insights shows a summary, aspect analysis with evidence sentences, filterable reviews with LSTM predictions, and the trust table.")
e += [rect(0.6, 1.35, 6.0, 4.75, fill=WHITE), img(ASSETS / "screens" / "01b_home_pillars.png", 0.8, 1.5, w=5.6),
      text(0.6, 5.2, 6.0, 0.55, [P("Home: the three pillars and the pipeline", 14, True, NAVY, align="c", space=0)], "middle"),
      rect(6.75, 1.35, 6.0, 4.75, fill=WHITE), img(ASSETS / "poster_figs" / "03b_insights_aspects_crop.png", 7.3, 1.5, w=4.9),
      text(6.75, 5.2, 6.0, 0.55, [P("Product Insights: aspects with supporting evidence", 14, True, NAVY, align="c", space=0)], "middle")]
e += callout(0.6, 6.25, 12.15, 0.65, "Every insight links to the reviews behind it, so users can check the evidence.", 16)

# =============================================================== 15 DASH 2
e = slide("The dashboard (2 of 2)", "Review Trust has the threshold slider and flag table. Ask ReviewIQ answers from one product's reviews and shows its evidence. The Deep Learning Lab has a live classifier comparing three models with attention highlighting.")
shots = [("03c_insights_trust_crop.png", "Review Trust: slider and flag table"), ("05_ask_crop.png", "Ask ReviewIQ: answer with evidence"), ("07d_dl_live_crop.png", "Deep Learning Lab: live classifier")]
for i, (f, cap) in enumerate(shots):
    x = 0.6 + i * 4.1
    e += [rect(x, 1.4, 3.95, 4.25, fill=WHITE), img(ASSETS / "poster_figs" / f, x + 0.08, 1.5, w=3.79),
          text(x, 4.5, 3.95, 0.6, [P(cap, 14, True, NAVY, align="c", space=0)], "middle")]
e += callout(0.6, 5.9, 12.15, 0.85, "Live classifier: the three models can disagree on a mixed review, which is exactly why aspect-level analysis exists.", 17)

# =============================================================== 16 PROBLEMS
e = slide("Problems we found and fixed", "These came from testing and inspection, and they shaped the final design. Mention that we checked results rather than trusting them.")
rows = [("Backwards trust features", "Raw length and rating flagged long, negative reviews. Switched to directional features and account volume."),
        ("Ungrounded answers", "A warranty question got a confident \"No\" at low similarity. Added the similarity floor."),
        ("Slow, dimmed pages", "VADER features for 50,000 reviews were rebuilt on every click (23 s). Cached."),
        ("Missing product images", "20 images used an unread field and the fallback link was blocked. Fixed the parser, built-in fallback."),
        ("Ignored sidebar clicks", "The radio's identity changed each page, so the first click was dropped. Stable key and callbacks."),
        ("Duplicate navigation", "Streamlit read a folder named pages as its own menu. Renamed it.")]
for i, (t, d) in enumerate(rows):
    col, row = i % 2, i // 2
    x, y = 0.6 + col * 6.15, 1.4 + row * 1.78
    e += [rect(x, y, 6.0, 1.62, fill=WHITE), rect(x, y + 0.15, 0.09, 1.32, fill=ACCENT, line=ACCENT, radius=False),
          text(x + 0.22, y + 0.08, 5.7, 0.5, [P(t, 18, True, NAVY, space=0)], "middle"),
          text(x + 0.22, y + 0.6, 5.7, 1.0, [P(d, 14, False, MUTED, space=0)])]

# =============================================================== 17 LIMITS + FUTURE
e = slide("Limitations and future scope", "State the limitations plainly. They are real and they affect how the numbers should be read.")
e += card(0.6, 1.4, 6.0, 5.3, "Limitations", [
    P("Labels come from star ratings, so some are noisy", 17, bullet=True),
    P("Sample is a reviewer-grouped block of one English category", 17, bullet=True),
    P("DistilBERT saw one-sixth of the LSTM's data and no class weights", 17, bullet=True),
    P("Aspects use keyword lexicons: implicit aspects and sarcasm are missed", 17, bullet=True),
    P("Flagged reviews are anomalies, not confirmed fakes", 17, bullet=True),
    P("Answer faithfulness was not measured", 17, bullet=True)], accent=ACCENT)
e += card(6.75, 1.4, 6.0, 5.3, "Future scope", [
    P("Train DistilBERT on the full data with matched class weights", 17, bullet=True),
    P("Learned aspect extraction instead of keyword lists", 17, bullet=True),
    P("Human-annotated neutral and suspicious subsets", 17, bullet=True),
    P("Random sampling across several categories and languages", 17, bullet=True),
    P("Evaluate answer faithfulness, try larger generators", 17, bullet=True),
    P("Seller view that tracks complaint areas over time", 17, bullet=True)], accent=TEAL)

# =============================================================== 18 CONCLUSION
e = slide("Conclusion", "Summarise the three headline findings and what the system delivers.")
e += stat(0.6, 1.4, 3.95, 1.9, "0.718", "best macro-F1 (BiLSTM), 88.8% accuracy", ACCENT, SOFT_ORG)
e += stat(4.7, 1.4, 3.95, 1.9, "57x", "faster than DistilBERT on CPU", BLUE, SOFT_BLUE)
e += stat(8.8, 1.4, 3.95, 1.9, "20.2%", "of battery mentions are negative: top complaint", TEAL, SOFT_TEAL)
e += [text(0.6, 3.6, 12.1, 3.1, [
    P("NLP, deep learning and advanced ML on **one corpus** give buyers evidence-linked insight and sellers a ranked list of problems.", 19, bullet=True, space=8),
    P("A BiLSTM trained from scratch **matches DistilBERT's accuracy** once class weights are removed and is far cheaper to run.", 19, bullet=True, space=8),
    P("A **similarity floor** makes question answering decline unsupported questions instead of guessing.", 19, bullet=True, space=8),
    P("Results, limits and fixes are reported **honestly**, with confidence intervals and an ablation.", 19, bullet=True)])]

# =============================================================== 19 THANK YOU
e = slide(dark=True, notes="Thank the audience and invite questions.")
e += [rect(0, 0, W, H, fill=NAVY, line=NAVY, radius=False), rect(0, 4.9, W, 0.09, fill=ACCENT, line=ACCENT, radius=False)]
e += [text(0.7, 1.2, 11.9, 1.4, [P("Thank you", 64, True, WHITE, space=0)]),
      text(0.7, 2.7, 11.9, 0.9, [P("Questions?", 34, False, "FFB48F", space=0)]),
      text(0.7, 3.6, 11.9, 0.6, [P("ReviewIQ: AI-Based Product Review Intelligence and Consumer Insight System", 20, False, "DBE3F3", space=0)])]
for i, (n, r) in enumerate(TEAM):
    e += [text(0.7, 5.3 + i * 0.4, 3.6, 0.4, [P(n, 18, True, WHITE, space=0)], "middle"), text(4.2, 5.3 + i * 0.4, 2.6, 0.4, [P(r, 18, False, "C8D2E6", space=0)], "middle")]
e += [text(8.2, 5.3, 4.8, 0.5, [P(GUIDE, 22, True, WHITE, space=0)], "middle"), text(8.2, 5.8, 4.8, 0.4, [P("Project Guide", 15, False, "C8D2E6", space=0)]),
      text(8.2, 6.3, 4.8, 0.4, [P(COLLEGE, 16, False, "C8D2E6", space=0)])]


# ============================================================ RENDER: PPTX
def rgb(h):
    return RGBColor.from_string(h)


def add_runs(par, s, size, bold, color):
    parts = re.split(r"(\*\*.+?\*\*)", s)
    for part in parts:
        if not part:
            continue
        b = part.startswith("**") and part.endswith("**")
        r = par.add_run()
        r.text = part[2:-2] if b else part
        r.font.size = Pt(size)
        r.font.bold = bool(bold or b)
        r.font.name = FONT
        r.font.color.rgb = rgb(color)


def set_bullet(par, size):
    pPr = par._p.get_or_add_pPr()
    pPr.set("marL", str(int(Pt(size) * 1.1)))
    pPr.set("indent", str(-int(Pt(size) * 1.1)))
    for tag in ("a:buNone", "a:buChar"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", "•")


def build_pptx(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    n = len(SLIDES)
    for idx, s in enumerate(SLIDES, 1):
        sl = prs.slides.add_slide(blank)
        if not s["dark"]:
            bg = sl.background.fill
            bg.solid()
            bg.fore_color.rgb = rgb(PAPER)
            # title + chrome
            bar = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.42), Inches(0.12), Inches(0.62))
            bar.fill.solid(); bar.fill.fore_color.rgb = rgb(ACCENT); bar.line.fill.background()
            tb = sl.shapes.add_textbox(Inches(0.85), Inches(0.3), Inches(11.8), Inches(0.85))
            tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            add_runs(tf.paragraphs[0], s["title"], 30, True, NAVY)
            ft = sl.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(9), Inches(0.35))
            add_runs(ft.text_frame.paragraphs[0], f"ReviewIQ  |  {COLLEGE}", 11, False, MUTED)
            fn = sl.shapes.add_textbox(Inches(11.7), Inches(7.05), Inches(1.1), Inches(0.35))
            fn.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
            add_runs(fn.text_frame.paragraphs[0], f"{idx} / {n}", 11, False, MUTED)
        for el in s["els"]:
            t = el["t"]
            if t == "rect":
                shp = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if el["radius"] else MSO_SHAPE.RECTANGLE,
                                          Inches(el["x"]), Inches(el["y"]), Inches(el["w"]), Inches(el["h"]))
                if el["radius"]:
                    shp.adjustments[0] = 0.05
                shp.fill.solid(); shp.fill.fore_color.rgb = rgb(el["fill"])
                if el["line"] == el["fill"]:
                    shp.line.fill.background()
                else:
                    shp.line.color.rgb = rgb(el["line"]); shp.line.width = Pt(el["lw"])
                shp.shadow.inherit = False
            elif t == "arrow":
                shp = sl.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(el["x"]), Inches(el["y"]), Inches(el["w"]), Inches(el["h"]))
                shp.fill.solid(); shp.fill.fore_color.rgb = rgb(el["fill"]); shp.line.fill.background(); shp.shadow.inherit = False
            elif t == "img":
                im = Image.open(el["path"])
                w = el["w"]; h = el["h"] if el["h"] else w * im.height / im.width
                sl.shapes.add_picture(el["path"], Inches(el["x"]), Inches(el["y"]), Inches(w), Inches(h))
            elif t == "text":
                tb = sl.shapes.add_textbox(Inches(el["x"]), Inches(el["y"]), Inches(el["w"]), Inches(el["h"]))
                tf = tb.text_frame; tf.word_wrap = True
                tf.margin_left = tf.margin_right = Inches(el["pad"]); tf.margin_top = tf.margin_bottom = Inches(0.03)
                tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE}[el["anchor"]]
                for i, pr in enumerate(el["paras"]):
                    par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    par.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[pr["align"]]
                    par.space_after = Pt(pr["space"])
                    add_runs(par, pr["s"], pr["size"], pr["bold"], pr["color"])
                    if pr["bullet"]:
                        set_bullet(par, pr["size"])
            elif t == "table":
                rows, cols = len(el["rows"]), len(el["colw"])
                gt = sl.shapes.add_table(rows, cols, Inches(el["x"]), Inches(el["y"]), Inches(el["w"]), Inches(el["rowh"] * rows))
                tbl = gt.table
                for c, cw in enumerate(el["colw"]):
                    tbl.columns[c].width = Inches(cw)
                for r in range(rows):
                    tbl.rows[r].height = Inches(el["rowh"])
                    for c in range(cols):
                        cell = tbl.cell(r, c)
                        cell.margin_left = cell.margin_right = Inches(0.1)
                        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = rgb(NAVY if r == 0 else (SOFT_ORG if r in el["hl"] else WHITE))
                        par = cell.text_frame.paragraphs[0]
                        par.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT
                        isbest = (r, c) in el["best"]
                        add_runs(par, el["rows"][r][c], el["size"], r == 0 or isbest, WHITE if r == 0 else (ACCENT if isbest else INK))
        if s["notes"]:
            sl.notes_slide.notes_text_frame.text = s["notes"]
    prs.core_properties.title = "ReviewIQ: AI-Based Product Review Intelligence and Consumer Insight System"
    prs.core_properties.author = ", ".join(n for n, _ in TEAM)
    prs.core_properties.subject = f"Project presentation, {COLLEGE}. Guide: {GUIDE}"
    prs.save(path)


# ============================================================ RENDER: HTML preview
def esc_runs(s):
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", html.escape(s, quote=False))


def build_html(path):
    css = f"""
    body {{ margin:0; background:#888; font-family:'DejaVu Sans',Arial,sans-serif; }}
    .slide {{ position:relative; width:{W*96:.0f}px; height:{H*96:.0f}px; background:#{PAPER}; overflow:hidden; margin:0 0 20px 0; }}
    .abs {{ position:absolute; box-sizing:border-box; }}
    td {{ box-sizing:border-box; overflow:hidden; white-space:nowrap; }}
    .txt {{ overflow:hidden; line-height:1.2; }}
    .p {{ position:relative; }}
    .p.bul {{ padding-left:1.1em; }} .p.bul::before {{ content:'\\2022'; position:absolute; left:0; }}
    """
    out = [f"<html><head><meta charset='utf-8'><style>{css}</style></head><body>"]
    n = len(SLIDES)
    for idx, s in enumerate(SLIDES, 1):
        out.append(f"<div class='slide' id='s{idx}' style='background:#{NAVY if s['dark'] else PAPER}'>")
        if not s["dark"]:
            out.append(f"<div class='abs' style='left:{0.6*96}px;top:{0.42*96}px;width:{0.12*96}px;height:{0.62*96}px;background:#{ACCENT}'></div>")
            out.append(f"<div class='abs txt chk' style='left:{0.85*96}px;top:{0.3*96}px;width:{11.8*96}px;height:{0.85*96}px;font-size:{30*96/72}px;font-weight:700;color:#{NAVY};display:flex;align-items:center'>{html.escape(s['title'])}</div>")
            out.append(f"<div class='abs' style='left:{0.6*96}px;top:{7.05*96}px;font-size:{11*96/72}px;color:#{MUTED}'>ReviewIQ  |  {COLLEGE}</div>")
            out.append(f"<div class='abs' style='right:{(W-12.8)*96}px;top:{7.05*96}px;font-size:{11*96/72}px;color:#{MUTED}'>{idx} / {n}</div>")
        for el in s["els"]:
            t = el["t"]
            geo = f"left:{el['x']*96}px;top:{el['y']*96}px;width:{el['w']*96}px;"
            if t == "rect":
                rad = "border-radius:10px;" if el["radius"] else ""
                out.append(f"<div class='abs' style='{geo}height:{el['h']*96}px;background:#{el['fill']};border:{el['lw']}px solid #{el['line']};{rad}'></div>")
            elif t == "arrow":
                out.append(f"<div class='abs' style='{geo}height:{el['h']*96}px;background:#{el['fill']};clip-path:polygon(0 25%,60% 25%,60% 0,100% 50%,60% 100%,60% 75%,0 75%)'></div>")
            elif t == "img":
                im = Image.open(el["path"]); h = el["h"] if el["h"] else el["w"] * im.height / im.width
                out.append(f"<img class='abs' src='file://{el['path']}' style='{geo}height:{h*96}px'>")
            elif t == "text":
                jc = {"top": "flex-start", "middle": "center"}[el["anchor"]]
                inner = []
                for pr in el["paras"]:
                    cls = "p bul" if pr["bullet"] else "p"
                    al = {"l": "left", "c": "center", "r": "right"}[pr["align"]]
                    inner.append(f"<div class='{cls}' style='font-size:{pr['size']*96/72}px;font-weight:{700 if pr['bold'] else 400};color:#{pr['color']};text-align:{al};margin-bottom:{pr['space']*96/72}px'>{esc_runs(pr['s'])}</div>")
                out.append(f"<div class='abs txt chk' style='{geo}height:{el['h']*96}px;padding:3px {el['pad']*96}px;display:flex;flex-direction:column;justify-content:{jc}'><div>{''.join(inner)}</div></div>")
            elif t == "table":
                rows = []
                for r, row in enumerate(el["rows"]):
                    tds = []
                    for c, cell in enumerate(row):
                        best = (r, c) in el["best"]
                        bg = NAVY if r == 0 else (SOFT_ORG if r in el["hl"] else WHITE)
                        col = WHITE if r == 0 else (ACCENT if best else INK)
                        fw = 700 if (r == 0 or best) else 400
                        tds.append(f"<td style='width:{el['colw'][c]*96}px;height:{el['rowh']*96}px;padding:0 10px;background:#{bg};color:#{col};font-weight:{fw};text-align:{'left' if c == 0 else 'right'};font-size:{el['size']*96/72}px;border:1px solid #{LINE}'>{html.escape(cell)}</td>")
                    rows.append("<tr>" + "".join(tds) + "</tr>")
                out.append(f"<table class='abs' style='{geo}border-collapse:collapse;table-layout:fixed'>{''.join(rows)}</table>")
        out.append("</div>")
    out.append("</body></html>")
    pathlib.Path(path).write_text("\n".join(out))


if __name__ == "__main__":
    build_pptx(sys.argv[1])
    build_html(sys.argv[2])
    print(f"{len(SLIDES)} slides written")
